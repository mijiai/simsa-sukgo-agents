"""report_generate_service end-to-end 검증 — 새 파이프라인."""

import json
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document

from src.agents.report.schemas import ReportRequest
from src.agents.report.service import report_generate_service
from src.agents.report.template_spec import load_template
from src.common.constants import RiskLevel
from src.storage.schemas import AgentName, JobStatus


def setup_module() -> None:
    load_template.cache_clear()


def _raw_payload() -> dict:
    return {
        "company_name": "ACME",
        "company_id": "c-1",
        "news": [],
        "uploaded_files": [],
        "extracted_tables": [],
        "extracted_images": [],
    }


def _result_payload(*, section_insights: list[dict] | None = None) -> dict:
    return {
        "job_id": "job-1",
        "company_name": "ACME",
        "company_id": "c-1",
        "risk_level": "MEDIUM",
        "risk_score": 55.0,
        "summary": "분석 요약",
        "key_risk_factors": ["요인1"],
        "positive_signals": ["신호1"],
        "data_gaps": ["재무제표 미확보"],
        "section_insights": section_insights or [],
    }


def _make_deps(
    *,
    raw: dict | None = None,
    result: dict | None = None,
    plan_dict: dict | None = None,
):
    raw_bytes = json.dumps(raw or _raw_payload(), ensure_ascii=False).encode("utf-8")
    result_bytes = json.dumps(result or _result_payload(), ensure_ascii=False).encode("utf-8")

    blob = MagicMock()
    blob.download = AsyncMock(side_effect=[raw_bytes, result_bytes])
    blob.upload = AsyncMock()
    blob.generate_sas_url = MagicMock(side_effect=lambda path, _td: f"https://sas/{path}")

    tables = MagicMock()
    tables.jobs = MagicMock()
    tables.jobs.update_status = AsyncMock()
    tables.agent_status = MagicMock()
    tables.agent_status.update_running = AsyncMock()
    tables.agent_status.update_done = AsyncMock()
    tables.agent_status.update_failed = AsyncMock()
    tables.jobs_ref = MagicMock()
    tables.jobs_ref.update_after_done = AsyncMock()

    anthropic = MagicMock()
    anthropic.model = "claude-sonnet-4-6"
    anthropic.complete_json = AsyncMock(return_value=plan_dict or {"sections": {}})

    return blob, tables, anthropic


# ───────────────────── happy path ─────────────────────


async def test_service_uploads_both_md_and_docx_and_returns_both_urls() -> None:
    blob, tables, anthropic = _make_deps()
    request = ReportRequest(job_id="job-1")

    response = await report_generate_service(request, blob, tables, anthropic, sas_expiry_hours=168)

    # 두 번 upload (docx + md)
    assert blob.upload.await_count == 2
    paths_uploaded = [c.args[0] for c in blob.upload.call_args_list]
    assert "jobs/job-1/report/report.docx" in paths_uploaded
    assert "jobs/job-1/report/report.md" in paths_uploaded

    docx_call = next(c for c in blob.upload.call_args_list if c.args[0].endswith(".docx"))
    assert (
        docx_call.kwargs["content_type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    md_call = next(c for c in blob.upload.call_args_list if c.args[0].endswith(".md"))
    assert md_call.kwargs["content_type"] == "text/markdown; charset=utf-8"

    # docx 바이트가 valid python-docx 로 다시 열림
    docx_bytes = docx_call.args[1]
    Document(BytesIO(docx_bytes))  # raise 없이 통과해야 함

    # SAS URL 두 개 발급
    assert blob.generate_sas_url.call_count == 2
    assert response.report_url == "https://sas/jobs/job-1/report/report.md"
    assert response.docx_url == "https://sas/jobs/job-1/report/report.docx"
    assert response.report_blob_path == "jobs/job-1/report/report.md"
    assert response.docx_blob_path == "jobs/job-1/report/report.docx"

    assert response.status == "done"
    assert response.risk_level is RiskLevel.MEDIUM
    assert response.risk_score == 55.0


async def test_service_calls_llm_exactly_twice_planner_then_narrative() -> None:
    """LLM 호출 횟수 정책 — planner 1회 + narrative writer 1회 = 총 2회.

    PR-1 (narrative writer 도입) 이후 정책 갱신. planner 만 1회였던 것은
    Option B 의 1단계, narrative writer 가 2단계.
    """
    blob, tables, anthropic = _make_deps()
    request = ReportRequest(job_id="job-2")
    await report_generate_service(request, blob, tables, anthropic, sas_expiry_hours=168)
    assert anthropic.complete_json.await_count == 2


async def test_service_propagates_section_insights_into_docx_bullets() -> None:
    """financial.section_insights 의 bullet 이 docx 표 아래에 그대로 포함되는지."""
    insights = [
        {
            "section_id": "5_finance",
            "table_id": "재무제표",
            "bullets": ["자산총계 1200 (전년 +200)", "부채비율 50%"],
            "cited_data_points": ["자산총계 1200"],
        }
    ]
    plan_dict = {
        "sections": {
            "5_finance": {
                "tables": {
                    "재무제표": {
                        "data_gap": False,
                        "columns": ["구분", "2023", "2024"],
                        "rows": [["자산총계", 1000, 1200]],
                    }
                },
                "images": {},
            }
        }
    }
    blob, tables, anthropic = _make_deps(
        result=_result_payload(section_insights=insights),
        plan_dict=plan_dict,
    )
    request = ReportRequest(job_id="job-3")
    await report_generate_service(request, blob, tables, anthropic, sas_expiry_hours=168)

    docx_call = next(c for c in blob.upload.call_args_list if c.args[0].endswith(".docx"))
    docx_bytes = docx_call.args[1]
    reopened = Document(BytesIO(docx_bytes))
    text = "\n".join(p.text for p in reopened.paragraphs)
    assert "자산총계 1200 (전년 +200)" in text
    assert "부채비율 50%" in text


# ───────────────────── status / jobs_ref ─────────────────────


async def test_service_updates_status_and_jobs_ref() -> None:
    blob, tables, anthropic = _make_deps()
    request = ReportRequest(job_id="job-4")

    await report_generate_service(request, blob, tables, anthropic, sas_expiry_hours=168)

    # AnalysisJobs.status: REPORTING → DONE
    statuses = [c.args[1] for c in tables.jobs.update_status.call_args_list]
    assert statuses[0] is JobStatus.REPORTING
    assert statuses[-1] is JobStatus.DONE

    tables.agent_status.update_done.assert_awaited_once_with(
        "job-4", AgentName.REPORT, output_blob_path="jobs/job-4/report/report.docx"
    )

    tables.jobs_ref.update_after_done.assert_awaited_once()
    ref_kwargs = tables.jobs_ref.update_after_done.call_args.kwargs
    assert ref_kwargs["company_id"] == "c-1"
    assert ref_kwargs["risk_level"] is RiskLevel.MEDIUM


async def test_service_skips_jobs_ref_when_no_company_id() -> None:
    raw = _raw_payload()
    raw["company_id"] = None
    result = _result_payload()
    result["company_id"] = None
    blob, tables, anthropic = _make_deps(raw=raw, result=result)

    request = ReportRequest(job_id="job-5")
    await report_generate_service(request, blob, tables, anthropic, sas_expiry_hours=168)

    tables.jobs_ref.update_after_done.assert_not_called()


# ───────────────────── failures ─────────────────────


async def test_service_blob_download_failure_marks_failed() -> None:
    blob, tables, anthropic = _make_deps()
    blob.download = AsyncMock(side_effect=RuntimeError("blob read boom"))

    request = ReportRequest(job_id="job-6")
    with pytest.raises(RuntimeError):
        await report_generate_service(request, blob, tables, anthropic, sas_expiry_hours=168)

    tables.agent_status.update_failed.assert_awaited_once()
    anthropic.complete_json.assert_not_called()
    blob.upload.assert_not_called()


async def test_service_invalid_risk_level_marks_failed() -> None:
    bad_result = _result_payload()
    bad_result["risk_level"] = "INVALID"
    blob, tables, anthropic = _make_deps(result=bad_result)

    request = ReportRequest(job_id="job-7")
    with pytest.raises(ValueError):
        await report_generate_service(request, blob, tables, anthropic, sas_expiry_hours=168)

    tables.agent_status.update_failed.assert_awaited_once()


async def test_service_planner_llm_failure_still_produces_docx_via_fallback() -> None:
    """planner LLM 이 실패해도 fallback plan 으로 보고서는 생성되어야 함 (모든 슬롯 data_gap)."""
    blob, tables, anthropic = _make_deps()
    anthropic.complete_json = AsyncMock(side_effect=RuntimeError("LLM down"))

    request = ReportRequest(job_id="job-8")
    response = await report_generate_service(request, blob, tables, anthropic, sas_expiry_hours=168)

    # 보고서는 정상 생성됨 (단, 표는 [자료 미확보] placeholder)
    assert response.status == "done"
    assert blob.upload.await_count == 2
    docx_call = next(c for c in blob.upload.call_args_list if c.args[0].endswith(".docx"))
    docx_bytes = docx_call.args[1]
    reopened = Document(BytesIO(docx_bytes))
    text = "\n".join(p.text for p in reopened.paragraphs)
    assert "[자료 미확보]" in text


async def test_service_sas_expiry_hours_propagates_to_both_urls() -> None:
    blob, tables, anthropic = _make_deps()
    request = ReportRequest(job_id="job-9")
    await report_generate_service(request, blob, tables, anthropic, sas_expiry_hours=24)

    # 두 URL 모두 24시간 timedelta 로 호출됨
    sas_calls = blob.generate_sas_url.call_args_list
    assert len(sas_calls) == 2
    for c in sas_calls:
        assert c.args[1].total_seconds() == 24 * 3600
