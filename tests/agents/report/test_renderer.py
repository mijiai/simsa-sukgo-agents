"""Renderer 검증 — docx 바이트 round-trip + markdown 포맷."""

from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

from docx import Document
from PIL import Image

from src.agents.report.renderer import render_report_docx, render_report_markdown
from src.agents.report.sections import ImageContent, SectionContent, TableContent
from src.agents.report.template_spec import load_template
from src.common.constants import ReportSection


def setup_module() -> None:
    load_template.cache_clear()


def _section_with_table(
    section_id: ReportSection = ReportSection.FINANCE,
    *,
    title: str = "재무 현황",
    table_data_gap: bool = False,
    bullets: list[str] | None = None,
) -> SectionContent:
    return SectionContent(
        section_id=section_id,
        number="5",
        title=title,
        tables=[
            TableContent(
                table_id="재무제표",
                title="주요 재무제표 추이",
                columns=["구분", "2023", "2024"],
                rows=[["자산총계", 1000, 1200], ["부채총계", 500, 600]],
                source_label="※ 출처: 외부감사 회계감사보고서",
                bullets=bullets or ["자산 +200 (전년 대비)"],
                data_gap=table_data_gap,
            )
        ],
    )


def _make_png() -> bytes:
    img = Image.new("RGB", (200, 100), color="green")
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# ───────────────────── DOCX ─────────────────────


async def test_render_docx_produces_valid_zip_openable_by_python_docx() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.FINANCE: _section_with_table()}
    docx_bytes = await render_report_docx(template, sections, blob=None)

    # round-trip: python-docx 가 다시 열 수 있어야 함
    reopened = Document(BytesIO(docx_bytes))
    headings = [p.text for p in reopened.paragraphs if p.style.name.startswith("Heading")]
    # 9개 섹션 헤딩 모두 포함
    assert any("5. 재무 현황" in h for h in headings)
    assert any("9. 종합 의견" in h for h in headings)


async def test_render_docx_includes_table_with_correct_cells() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.FINANCE: _section_with_table()}
    docx_bytes = await render_report_docx(template, sections, blob=None)
    reopened = Document(BytesIO(docx_bytes))

    found = False
    for tbl in reopened.tables:
        cell_texts = [tbl.rows[0].cells[i].text for i in range(len(tbl.rows[0].cells))]
        if cell_texts == ["구분", "2023", "2024"]:
            found = True
            assert tbl.rows[1].cells[0].text == "자산총계"
            assert tbl.rows[1].cells[1].text == "1000"
            assert tbl.rows[1].cells[2].text == "1200"
            break
    assert found, "예상한 컬럼 헤더의 표를 찾지 못함"


async def test_render_docx_data_gap_shows_placeholder() -> None:
    template = load_template("loan_application_v1")
    sections = {
        ReportSection.FINANCE: _section_with_table(table_data_gap=True),
    }
    docx_bytes = await render_report_docx(template, sections, blob=None)
    reopened = Document(BytesIO(docx_bytes))
    text = "\n".join(p.text for p in reopened.paragraphs)
    assert "[자료 미확보]" in text


async def test_render_docx_image_data_gap_when_blob_none() -> None:
    template = load_template("loan_application_v1")
    sections = {
        ReportSection.AFFILIATES: SectionContent(
            section_id=ReportSection.AFFILIATES,
            number="8",
            title="관계사 현황",
            images=[
                ImageContent(
                    image_id="지분 구조도",
                    title="지분 구조도",
                    blob_path="jobs/J/input/ownership.png",
                    suspected_role="ownership_chart",
                    data_gap=False,
                )
            ],
        ),
    }
    # blob=None → placeholder 처리
    docx_bytes = await render_report_docx(template, sections, blob=None)
    reopened = Document(BytesIO(docx_bytes))
    text = "\n".join(p.text for p in reopened.paragraphs)
    assert "지분 구조도" in text
    assert "[자료 미확보]" in text


async def test_render_docx_downloads_image_when_blob_present() -> None:
    template = load_template("loan_application_v1")
    sections = {
        ReportSection.AFFILIATES: SectionContent(
            section_id=ReportSection.AFFILIATES,
            number="8",
            title="관계사 현황",
            images=[
                ImageContent(
                    image_id="지분 구조도",
                    title="지분 구조도",
                    blob_path="jobs/J/input/ownership.png",
                    suspected_role="ownership_chart",
                )
            ],
        ),
    }
    blob = MagicMock()
    blob.download = AsyncMock(return_value=_make_png())
    docx_bytes = await render_report_docx(template, sections, blob=blob)
    blob.download.assert_awaited_once_with("jobs/J/input/ownership.png")
    # round-trip 으로 inline shape 1개 이상 존재 확인
    reopened = Document(BytesIO(docx_bytes))
    shape_count = len(reopened.inline_shapes)
    assert shape_count >= 1


async def test_render_docx_image_download_failure_falls_back_to_placeholder() -> None:
    template = load_template("loan_application_v1")
    sections = {
        ReportSection.AFFILIATES: SectionContent(
            section_id=ReportSection.AFFILIATES,
            number="8",
            title="관계사 현황",
            images=[
                ImageContent(
                    image_id="지분 구조도",
                    title="지분 구조도",
                    blob_path="jobs/J/input/ownership.png",
                    suspected_role="ownership_chart",
                )
            ],
        ),
    }
    blob = MagicMock()
    blob.download = AsyncMock(side_effect=RuntimeError("blob read boom"))
    docx_bytes = await render_report_docx(template, sections, blob=blob)
    reopened = Document(BytesIO(docx_bytes))
    text = "\n".join(p.text for p in reopened.paragraphs)
    assert "이미지 로드 실패" in text


# ───────────────────── Markdown ─────────────────────


def test_render_markdown_includes_all_section_headings() -> None:
    template = load_template("loan_application_v1")
    sections = {}  # 모두 missing → placeholder
    md = render_report_markdown(template, sections)
    for spec in template.sections:
        assert f"# {spec.number}. {spec.title}" in md


def test_render_markdown_table_uses_gfm_pipe_format() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.FINANCE: _section_with_table()}
    md = render_report_markdown(template, sections)
    assert "**주요 재무제표 추이**" in md
    assert "| 구분 | 2023 | 2024 |" in md
    assert "| --- | --- | --- |" in md
    assert "| 자산총계 | 1000 | 1200 |" in md
    assert "- 자산 +200 (전년 대비)" in md


def test_render_markdown_data_gap_shows_italic_placeholder() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.FINANCE: _section_with_table(table_data_gap=True)}
    md = render_report_markdown(template, sections)
    assert "_[자료 미확보]" in md


def test_render_markdown_section_bullets_appended() -> None:
    template = load_template("loan_application_v1")
    sections = {
        ReportSection.CONCLUSION: SectionContent(
            section_id=ReportSection.CONCLUSION,
            number="9",
            title="종합 의견",
            section_bullets=["전체 위험 MEDIUM", "추가 자료 확인 필요"],
        )
    }
    md = render_report_markdown(template, sections)
    assert "- 전체 위험 MEDIUM" in md
    assert "- 추가 자료 확인 필요" in md


# ───────────────────── narrative paragraph ─────────────────────


def test_render_markdown_narrative_appears_above_tables() -> None:
    """narrative 가 섹션 헤딩 다음 + 표보다 앞에 출력."""
    template = load_template("loan_application_v1")
    section = _section_with_table()
    section.narrative = "동사는 '24년 자산 1,200억원 기록함. 전년 대비 200억원 증가."
    sections = {ReportSection.FINANCE: section}
    md = render_report_markdown(template, sections)
    # heading → narrative → table 순서
    h_idx = md.index("# 5. 재무 현황")
    n_idx = md.index("동사는 '24년 자산")
    t_idx = md.index("**주요 재무제표 추이**")
    assert h_idx < n_idx < t_idx


def test_render_markdown_narrative_split_by_blank_line() -> None:
    """\\n\\n 구분 paragraph 가 markdown 에서도 별도 줄로 분리."""
    template = load_template("loan_application_v1")
    section = _section_with_table()
    section.narrative = "첫 paragraph 본문.\n\n둘째 paragraph 본문."
    sections = {ReportSection.FINANCE: section}
    md = render_report_markdown(template, sections)
    assert "첫 paragraph 본문." in md
    assert "둘째 paragraph 본문." in md


def test_render_markdown_no_narrative_block_when_empty() -> None:
    """narrative 가 빈 문자열이면 출력 자체 생략."""
    template = load_template("loan_application_v1")
    section = _section_with_table()
    # narrative default = ""
    sections = {ReportSection.FINANCE: section}
    md = render_report_markdown(template, sections)
    # heading 바로 다음에 빈 줄 / table 만 있어야 함 (유사 본문 텍스트 없음)
    after_heading = md.split("# 5. 재무 현황", 1)[1].split("**주요 재무제표 추이**", 1)[0]
    # 빈 문자열 또는 공백/개행만 — 실제 prose 가 끼어있지 않은지 확인
    assert "본문" not in after_heading
    assert after_heading.strip() == ""


async def test_render_docx_includes_narrative_paragraph() -> None:
    """docx 본문에 narrative 가 paragraph 로 들어가는지 (round-trip 확인)."""
    template = load_template("loan_application_v1")
    section = _section_with_table()
    section.narrative = "DOCX_NARRATIVE_MARKER 동사는 안정적인 수익성 보유 중."
    sections = {ReportSection.FINANCE: section}
    docx_bytes = await render_report_docx(template, sections)
    doc = Document(BytesIO(docx_bytes))
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "DOCX_NARRATIVE_MARKER" in body_text


async def test_render_docx_no_narrative_when_empty() -> None:
    template = load_template("loan_application_v1")
    section = _section_with_table()
    sections = {ReportSection.FINANCE: section}
    docx_bytes = await render_report_docx(template, sections)
    doc = Document(BytesIO(docx_bytes))
    body_text = "\n".join(p.text for p in doc.paragraphs)
    # 표 제목/bullet 외에 narrative 형 prose 가 없어야 함
    assert "DOCX_NARRATIVE_MARKER" not in body_text
