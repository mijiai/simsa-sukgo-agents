from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

from docx import Document

from src.agents.report.templates import (
    extract_docx_text,
    get_cached_templates,
    load_report_samples,
    reset_templates_for_tests,
)


def _make_docx(lines: list[str]) -> bytes:
    doc = Document()
    for line in lines:
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        else:
            doc.add_paragraph(line)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_docx_with_table(rows: list[list[str]]) -> bytes:
    doc = Document()
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r_idx, row in enumerate(rows):
        for c_idx, cell in enumerate(row):
            table.rows[r_idx].cells[c_idx].text = cell
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_docx_text_paragraphs() -> None:
    data = _make_docx(["# 기업 개요", "ACME Corp 는 IT 기업이다.", "본사 위치: 서울"])
    text = extract_docx_text(data)
    assert "기업 개요" in text
    assert "ACME Corp 는 IT 기업이다." in text
    assert "본사 위치: 서울" in text


def test_extract_docx_text_skips_empty_paragraphs() -> None:
    data = _make_docx(["내용 있음", "", "  ", "또 다른 내용"])
    text = extract_docx_text(data)
    lines = text.split("\n")
    assert "내용 있음" in lines
    assert "또 다른 내용" in lines
    assert "" not in lines


def test_extract_docx_text_includes_tables() -> None:
    data = _make_docx_with_table(
        [["항목", "2022", "2023"], ["매출", "100", "120"], ["순익", "10", "15"]]
    )
    text = extract_docx_text(data)
    assert "항목 | 2022 | 2023" in text
    assert "매출 | 100 | 120" in text


async def test_load_report_samples_caches_extracted_text() -> None:
    reset_templates_for_tests()

    sample1 = _make_docx(["# 기업 개요", "샘플1 본문"])
    sample2 = _make_docx(["# 기업 개요", "샘플2 본문"])

    blob = MagicMock()
    blob.list_prefix = AsyncMock(
        return_value=[
            "templates/report_samples/sample_1.docx",
            "templates/report_samples/sample_2.docx",
            "templates/report_samples/notes.txt",
        ]
    )
    blob.download = AsyncMock(side_effect=[sample1, sample2])

    samples = await load_report_samples(blob, "templates/report_samples/")

    assert len(samples) == 2
    assert "샘플1 본문" in samples[0]
    assert "샘플2 본문" in samples[1]
    assert get_cached_templates() == samples
    blob.download.assert_any_await("templates/report_samples/sample_1.docx")
    blob.download.assert_any_await("templates/report_samples/sample_2.docx")


async def test_load_report_samples_skips_corrupt_docx() -> None:
    reset_templates_for_tests()
    good = _make_docx(["좋은 샘플"])
    bad = b"not a real docx"

    blob = MagicMock()
    blob.list_prefix = AsyncMock(
        return_value=[
            "templates/report_samples/good.docx",
            "templates/report_samples/broken.docx",
        ]
    )
    blob.download = AsyncMock(side_effect=[good, bad])

    samples = await load_report_samples(blob, "templates/report_samples/")
    assert len(samples) == 1
    assert "좋은 샘플" in samples[0]


async def test_load_report_samples_empty_prefix_returns_empty_cache() -> None:
    reset_templates_for_tests(["initial"])

    blob = MagicMock()
    blob.list_prefix = AsyncMock(return_value=[])
    blob.download = AsyncMock()

    samples = await load_report_samples(blob, "templates/report_samples/")
    assert samples == []
    assert get_cached_templates() == []
    blob.download.assert_not_called()
