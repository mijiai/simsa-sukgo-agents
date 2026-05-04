"""PR4 폴리싱 검증 — base docx 상속 / 필수 이미지 정책 / 별첨 자동 분리."""

from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

from docx import Document
from docx.shared import Pt

from src.agents.report.renderer import (
    collect_appendix_tables,
    render_report_docx,
    render_report_markdown,
)
from src.agents.report.sections import ImageContent, SectionContent, TableContent
from src.agents.report.template_spec import load_template
from src.common.constants import ReportSection


def setup_module() -> None:
    load_template.cache_clear()


def _make_base_docx_with_marker(marker: str = "BASE_HEADER_MARKER") -> bytes:
    """\"base\" 역할의 docx — 본문에 식별 가능한 paragraph 가 있어야 본문 비움 검증 가능."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run(marker).font.size = Pt(11)
    # 스타일 1개 추가 — base 의 styles 가 보존되는지 검증용
    styles = doc.styles
    # custom paragraph style
    if "BasePolishMarker" not in [s.name for s in styles]:
        from docx.enum.style import WD_STYLE_TYPE

        style = styles.add_style("BasePolishMarker", WD_STYLE_TYPE.PARAGRAPH)
        style.font.size = Pt(15)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _section_with_long_table(row_count: int) -> SectionContent:
    rows = [[f"row{i}", i] for i in range(row_count)]
    return SectionContent(
        section_id=ReportSection.DEBT,
        number="6",
        title="차입금 현황",
        tables=[
            TableContent(
                table_id="차입금 명세",
                title="차입금 명세 (상세)",
                columns=["항목", "잔액"],
                rows=rows,
                source_label="※ 출처: 업무보고서",
                bullets=["차입금 의존도 상승"],
            )
        ],
    )


def _section_with_required_image(missing: bool = True) -> SectionContent:
    return SectionContent(
        section_id=ReportSection.AFFILIATES,
        number="8",
        title="관계사 현황",
        tables=[
            TableContent(
                table_id="관계사 목록",
                title="주요 관계사 현황",
                columns=["회사명"],
                rows=[["A 관계사"]],
            )
        ],
        images=[
            ImageContent(
                image_id="지분 구조도",
                title="지분 구조도",
                blob_path=None if missing else "jobs/J/input/x.png",
                suspected_role="ownership_chart",
                optional=False,  # ★ 필수 이미지
                data_gap=missing,
            )
        ],
    )


# ───────────────────── base docx 상속 ─────────────────────


async def test_render_docx_without_base_uses_blank_document() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.FINANCE: _section_with_long_table(5)}
    docx_bytes = await render_report_docx(template, sections, blob=None, base_docx_bytes=None)
    # base 없이도 정상 생성
    Document(BytesIO(docx_bytes))


async def test_render_docx_with_base_inherits_custom_style() -> None:
    base = _make_base_docx_with_marker()
    template = load_template("loan_application_v1")
    sections = {ReportSection.FINANCE: _section_with_long_table(5)}
    docx_bytes = await render_report_docx(template, sections, blob=None, base_docx_bytes=base)

    reopened = Document(BytesIO(docx_bytes))
    # base 의 custom style 이 보존됨
    style_names = [s.name for s in reopened.styles]
    assert "BasePolishMarker" in style_names


async def test_render_docx_with_base_clears_body_paragraphs() -> None:
    base = _make_base_docx_with_marker(marker="BASE_HEADER_MARKER_UNIQUE")
    template = load_template("loan_application_v1")
    sections = {ReportSection.DEBT: _section_with_long_table(5)}
    docx_bytes = await render_report_docx(template, sections, blob=None, base_docx_bytes=base)
    reopened = Document(BytesIO(docx_bytes))
    text = "\n".join(p.text for p in reopened.paragraphs)
    # base 본문은 비워짐 — marker 가 출력에 없어야 함
    assert "BASE_HEADER_MARKER_UNIQUE" not in text
    # 우리 컨텐츠는 정상 들어감
    assert "차입금 현황" in text
    assert "차입금 명세" in text


async def test_render_docx_with_corrupt_base_falls_back_to_blank() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.FINANCE: _section_with_long_table(5)}
    bad_base = b"this is not a valid docx"
    docx_bytes = await render_report_docx(template, sections, blob=None, base_docx_bytes=bad_base)
    # raise 없이 정상 생성
    Document(BytesIO(docx_bytes))


# ───────────────────── 필수 이미지 정책 ─────────────────────


async def test_render_docx_required_image_missing_omits_section_body() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.AFFILIATES: _section_with_required_image(missing=True)}
    docx_bytes = await render_report_docx(template, sections, blob=None)
    reopened = Document(BytesIO(docx_bytes))
    text = "\n".join(p.text for p in reopened.paragraphs)
    assert "[자료 미확보 — 추가 제출 필요]" in text
    # 표 본문은 출력되지 않아야 함
    table_texts = []
    for tbl in reopened.tables:
        for row in tbl.rows:
            table_texts.append(" ".join(c.text for c in row.cells))
    joined_table_text = "\n".join(table_texts)
    assert "A 관계사" not in joined_table_text


async def test_render_docx_required_image_present_keeps_section_body() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.AFFILIATES: _section_with_required_image(missing=False)}
    blob = MagicMock()
    # 이미지 다운로드는 더미로 실패하게 두지만 (PIL 검증 X), placeholder 가 나와도 표는 나와야 함
    blob.download = AsyncMock(side_effect=RuntimeError("image fetch fail"))
    docx_bytes = await render_report_docx(template, sections, blob=blob)
    reopened = Document(BytesIO(docx_bytes))
    table_texts = []
    for tbl in reopened.tables:
        for row in tbl.rows:
            table_texts.append(" ".join(c.text for c in row.cells))
    joined = "\n".join(table_texts)
    assert "A 관계사" in joined


def test_md_required_image_missing_shows_placeholder() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.AFFILIATES: _section_with_required_image(missing=True)}
    md = render_report_markdown(template, sections)
    assert "[자료 미확보 — 추가 제출 필요]" in md
    assert "A 관계사" not in md


# ───────────────────── 별첨 자동 분리 ─────────────────────


def test_collect_appendix_tables_picks_only_threshold_exceeders() -> None:
    sections = {
        ReportSection.DEBT: _section_with_long_table(50),
        ReportSection.FINANCE: _section_with_long_table(10),
    }
    appendix, ids = collect_appendix_tables(sections, row_threshold=35)
    assert len(appendix) == 1
    assert ids == {"차입금 명세"}
    assert appendix[0].table_id == "차입금 명세"
    assert len(appendix[0].rows) == 50


def test_collect_appendix_tables_skips_data_gap_tables() -> None:
    sections = {
        ReportSection.DEBT: SectionContent(
            section_id=ReportSection.DEBT,
            number="6",
            title="차입금 현황",
            tables=[
                TableContent(
                    table_id="missing",
                    title="missing",
                    columns=[],
                    rows=[],
                    data_gap=True,
                )
            ],
        )
    }
    appendix, ids = collect_appendix_tables(sections, row_threshold=1)
    assert appendix == []
    assert ids == set()


async def test_render_docx_long_table_moves_to_appendix_with_body_reference() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.DEBT: _section_with_long_table(50)}
    docx_bytes = await render_report_docx(template, sections, blob=None, appendix_row_threshold=35)
    reopened = Document(BytesIO(docx_bytes))

    # 본문에는 "별첨 참조" 안내, 표는 빠짐
    paragraphs_text = [p.text for p in reopened.paragraphs]
    assert any("별첨 참조" in t for t in paragraphs_text)

    # 별첨 헤딩이 추가됨
    headings = [p.text for p in reopened.paragraphs if p.style.name.startswith("Heading")]
    assert any("별첨" in h for h in headings)

    # 별첨 섹션의 표는 50행 그대로
    # docx 의 마지막 표가 별첨의 차입금 명세
    last_table = reopened.tables[-1]
    # header 1행 + 데이터 50행 = 51행
    assert len(last_table.rows) == 51


async def test_render_docx_long_table_below_threshold_stays_inline() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.DEBT: _section_with_long_table(20)}
    docx_bytes = await render_report_docx(template, sections, blob=None, appendix_row_threshold=35)
    reopened = Document(BytesIO(docx_bytes))
    paragraphs_text = [p.text for p in reopened.paragraphs]
    # 별첨 참조 메시지가 없어야 함
    assert not any("별첨 참조" in t for t in paragraphs_text)
    # 별첨 헤딩 없어야 함
    headings = [p.text for p in reopened.paragraphs if p.style.name.startswith("Heading")]
    assert not any("별첨" in h for h in headings)


def test_md_long_table_moves_to_appendix() -> None:
    template = load_template("loan_application_v1")
    sections = {ReportSection.DEBT: _section_with_long_table(50)}
    md = render_report_markdown(template, sections, appendix_row_threshold=35)
    assert "_(상세 표는 별첨 참조)_" in md
    # 별첨 헤딩 있음
    assert "# 별첨" in md
    # 본문에는 큰 표 없고 별첨에 있음 (row0 가 별첨 안에서만 나옴)
    appendix_idx = md.index("# 별첨")
    body_part = md[:appendix_idx]
    appendix_part = md[appendix_idx:]
    assert "row0 | 0" not in body_part
    assert "row0 | 0" in appendix_part
