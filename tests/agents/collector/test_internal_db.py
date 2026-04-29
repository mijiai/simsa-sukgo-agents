from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

import openpyxl
from docx import Document

from src.agents.collector.internal_db import (
    INTERNAL_PREFIX,
    get_company_data,
    load_internal_db,
    reset_for_tests,
)


def _make_xlsx(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "재무"
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_docx(lines: list[str]) -> bytes:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def test_load_internal_db_indexes_company_id_from_path() -> None:
    reset_for_tests()
    xlsx = _make_xlsx([["부채비율", 180], ["유동비율", 120]])

    blob = MagicMock()
    blob.list_prefix = AsyncMock(
        return_value=[
            "internal/companies/1248100998/financial.xlsx",
        ]
    )
    blob.download = AsyncMock(return_value=xlsx)

    db = await load_internal_db(blob)

    assert "1248100998" in db
    assert "부채비율 | 180" in db["1248100998"]
    assert get_company_data("1248100998") == db["1248100998"]
    assert get_company_data("missing-id") is None


async def test_load_internal_db_supports_multiple_companies_and_formats() -> None:
    reset_for_tests()
    xlsx = _make_xlsx([["매출", 1000]])
    docx = _make_docx(["DOCX 신용분석 본문"])

    blob = MagicMock()
    blob.list_prefix = AsyncMock(
        return_value=[
            "internal/companies/1111111111/financial.xlsx",
            "internal/companies/2222222222/financial.docx",
        ]
    )
    blob.download = AsyncMock(side_effect=[xlsx, docx])

    db = await load_internal_db(blob)

    assert set(db.keys()) == {"1111111111", "2222222222"}
    assert "매출 | 1000" in db["1111111111"]
    assert "DOCX 신용분석 본문" in db["2222222222"]


async def test_load_internal_db_skips_non_financial_filenames() -> None:
    reset_for_tests()

    blob = MagicMock()
    blob.list_prefix = AsyncMock(
        return_value=[
            "internal/companies/1111111111/notes.xlsx",  # wrong base name
            "internal/companies/2222222222/financial.txt",  # unsupported suffix
            "internal/_README.md",  # placeholder, not under a company_id
            "internal/companies/3333333333/financial.xlsx",  # OK
        ]
    )
    blob.download = AsyncMock(return_value=_make_xlsx([["x", 1]]))

    db = await load_internal_db(blob)

    assert set(db.keys()) == {"3333333333"}
    blob.download.assert_awaited_once_with("internal/companies/3333333333/financial.xlsx")


async def test_load_internal_db_skips_corrupt_file_continues_others() -> None:
    reset_for_tests()
    good = _make_xlsx([["good", 1]])
    bad = b"not a real xlsx"

    blob = MagicMock()
    blob.list_prefix = AsyncMock(
        return_value=[
            "internal/companies/aaaaaaaaaa/financial.xlsx",
            "internal/companies/bbbbbbbbbb/financial.xlsx",
        ]
    )
    blob.download = AsyncMock(side_effect=[bad, good])

    db = await load_internal_db(blob)

    assert "bbbbbbbbbb" in db
    assert "aaaaaaaaaa" not in db


async def test_load_internal_db_empty_prefix_resets_cache() -> None:
    reset_for_tests({"prev": "stale"})

    blob = MagicMock()
    blob.list_prefix = AsyncMock(return_value=[])
    blob.download = AsyncMock()

    db = await load_internal_db(blob)
    assert db == {}
    assert get_company_data("prev") is None
    blob.download.assert_not_called()


def test_internal_prefix_constant_is_companies_subdir() -> None:
    assert INTERNAL_PREFIX == "internal/companies/"
