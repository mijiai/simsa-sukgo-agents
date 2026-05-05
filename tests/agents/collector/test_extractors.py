from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

import openpyxl
import xlwt
from PIL import Image

from src.agents.collector.extractors import (
    PDF_MAX_TEXT_CHARS,
    SHEET_TEXT_MAX_CHARS,
    XLSX_MAX_ROWS_PER_SHEET,
    _is_sparse_sheets,
    extract_image,
    extract_pdf,
    extract_uploaded_file,
    extract_xls,
    extract_xlsx,
)


def _make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_xls(sheets: dict[str, list[list]]) -> bytes:
    wb = xlwt.Workbook()
    for name, rows in sheets.items():
        ws = wb.add_sheet(name)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                ws.write(r, c, val)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_png(width: int = 100, height: int = 50) -> bytes:
    img = Image.new("RGB", (width, height), color="red")
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _hand_rolled_pdf(text: str) -> bytes:
    """단일 페이지 텍스트 PDF (pdfplumber 가 추출 가능한 형태)."""
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content_stream = f"BT /F1 12 Tf 50 750 Td ({safe}) Tj ET".encode()
    objs: list[bytes] = []

    def add(o: bytes) -> None:
        objs.append(o)

    add(b"<< /Type /Catalog /Pages 2 0 R >>")
    add(b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    add(
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
    )
    add(
        b"<< /Length "
        + str(len(content_stream)).encode()
        + b" >>\nstream\n"
        + content_stream
        + b"\nendstream"
    )
    add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, obj in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF"
    ).encode()
    return bytes(out)


def _blob_with(data: bytes) -> MagicMock:
    blob = MagicMock()
    blob.download = AsyncMock(return_value=data)
    return blob


async def test_extract_xlsx_basic_separates_columns_and_rows() -> None:
    data = _make_xlsx(
        {
            "재무": [["구분", "2023", "2024"], ["자산총계", 1000, 1200]],
            "비율": [["지표", "값"], ["부채비율", 180.5]],
        }
    )
    sheets = await extract_xlsx(_blob_with(data), "any/path.xlsx")
    assert len(sheets) == 2

    by_name = {s["sheet_name"]: s for s in sheets}
    assert by_name["재무"]["columns"] == ["구분", "2023", "2024"]
    assert by_name["재무"]["rows"] == [["자산총계", 1000, 1200]]
    assert by_name["재무"]["truncated"] is False

    assert by_name["비율"]["columns"] == ["지표", "값"]
    assert by_name["비율"]["rows"] == [["부채비율", 180.5]]


async def test_extract_xlsx_truncates_rows_beyond_limit() -> None:
    rows = [["구분", "값"]] + [[f"row_{i}", i] for i in range(XLSX_MAX_ROWS_PER_SHEET + 50)]
    data = _make_xlsx({"S": rows})
    sheets = await extract_xlsx(_blob_with(data), "x.xlsx")
    assert len(sheets[0]["rows"]) == XLSX_MAX_ROWS_PER_SHEET
    assert sheets[0]["truncated"] is True


async def test_extract_xlsx_skips_empty_sheets() -> None:
    data = _make_xlsx({"empty": [], "real": [["A"], ["x"]]})
    sheets = await extract_xlsx(_blob_with(data), "x.xlsx")
    names = [s["sheet_name"] for s in sheets]
    assert "real" in names
    assert "empty" not in names


async def test_extract_xls_basic_separates_columns_and_rows() -> None:
    data = _make_xls(
        {
            "재무": [["구분", "2023", "2024"], ["자산총계", 1000, 1200]],
            "비율": [["지표", "값"], ["부채비율", 180.5]],
        }
    )
    sheets = await extract_xls(_blob_with(data), "any/path.xls")
    by_name = {s["sheet_name"]: s for s in sheets}
    assert by_name["재무"]["columns"] == ["구분", "2023", "2024"]
    assert by_name["재무"]["rows"] == [["자산총계", 1000, 1200]]
    assert by_name["재무"]["truncated"] is False
    assert by_name["비율"]["rows"] == [["부채비율", 180.5]]


async def test_extract_xls_truncates_rows_beyond_limit() -> None:
    rows = [["구분", "값"]] + [[f"row_{i}", i] for i in range(XLSX_MAX_ROWS_PER_SHEET + 50)]
    data = _make_xls({"S": rows})
    sheets = await extract_xls(_blob_with(data), "x.xls")
    assert len(sheets[0]["rows"]) == XLSX_MAX_ROWS_PER_SHEET
    assert sheets[0]["truncated"] is True


async def test_extract_uploaded_file_routes_xls_to_xls_extractor() -> None:
    """레거시 .xls 는 openpyxl 이 못 읽음 → xlrd 기반 extract_xls 로 가야 함.

    회귀 방지: 이 분기가 빠지면 .xls 업로드가 BadZipFile → 빈 raw.json → 보고서 [자료 미확보].
    """
    data = _make_xls({"재무": [["구분", "2024"], ["자산총계", 9999]]})
    result = await extract_uploaded_file(_blob_with(data), "p", "재무제표.xls")
    assert result is not None
    assert result["kind"] == "xlsx"  # downstream 분기 호환 (extracted_tables 로 들어감)
    assert result["content"][0]["sheet_name"] == "재무"
    assert result["content"][0]["rows"] == [["자산총계", 9999]]


async def test_extract_pdf_returns_text_and_page_count() -> None:
    data = _hand_rolled_pdf("PDF body content")
    result = await extract_pdf(_blob_with(data), "x.pdf")
    assert "PDF body content" in result["text"]
    assert result["page_count"] == 1
    assert result["truncated"] is False
    assert isinstance(result["tables"], list)


async def test_extract_pdf_truncates_text_at_limit() -> None:
    huge_pdf = _hand_rolled_pdf("X" * (PDF_MAX_TEXT_CHARS + 5000))
    result = await extract_pdf(_blob_with(huge_pdf), "x.pdf")
    assert len(result["text"]) <= PDF_MAX_TEXT_CHARS
    assert result["truncated"] is True


async def test_extract_image_returns_dimensions_and_role() -> None:
    blob = _blob_with(_make_png(width=480, height=273))
    result = await extract_image(blob, "jobs/J/input/ownership_chart.png", "ownership_chart.png")
    assert result["width"] == 480
    assert result["height"] == 273
    assert result["suspected_role"] == "ownership_chart"
    assert result["blob_path"] == "jobs/J/input/ownership_chart.png"
    assert result["original_filename"] == "ownership_chart.png"


async def test_extract_image_role_unknown_for_arbitrary_filename() -> None:
    blob = _blob_with(_make_png())
    result = await extract_image(blob, "p", "random.png")
    assert result["suspected_role"] == "unknown"


async def test_extract_image_role_product_catalog_keyword() -> None:
    blob = _blob_with(_make_png())
    result = await extract_image(blob, "p", "product_catalog_main.png")
    assert result["suspected_role"] == "product_catalog"


async def test_extract_uploaded_file_dispatches_by_extension() -> None:
    xlsx_blob = _blob_with(_make_xlsx({"S": [["a"], ["b"]]}))
    pdf_blob = _blob_with(_hand_rolled_pdf("text"))
    img_blob = _blob_with(_make_png())

    assert (await extract_uploaded_file(xlsx_blob, "p", "report.xlsx"))["kind"] == "xlsx"
    assert (await extract_uploaded_file(pdf_blob, "p", "report.pdf"))["kind"] == "pdf"
    assert (await extract_uploaded_file(img_blob, "p", "img.PNG"))["kind"] == "image"


async def test_extract_uploaded_file_returns_none_for_unsupported() -> None:
    blob = MagicMock()
    blob.download = AsyncMock()
    assert await extract_uploaded_file(blob, "p", "data.csv") is None
    assert await extract_uploaded_file(blob, "p", "notes.txt") is None
    blob.download.assert_not_called()


# ─── Sparse 시트 감지 + 텍스트 fallback (PDF 인쇄용 .xls 대응) ───


def test_is_sparse_sheets_empty_input() -> None:
    assert _is_sparse_sheets([]) is True


def test_is_sparse_sheets_clean_sheet_returns_false() -> None:
    sheets = [{"columns": ["구분", "2023", "2024"], "rows": [["자산", 100, 200]]}]
    assert _is_sparse_sheets(sheets) is False


def test_is_sparse_sheets_all_empty_columns_returns_true() -> None:
    sheets = [{"columns": ["", "", "", "", "", "", "", "", "", ""], "rows": []}]
    assert _is_sparse_sheets(sheets) is True


def test_is_sparse_sheets_mixed_70pct_threshold() -> None:
    # 7개 빈 / 3개 의미 → 70% 빈 → sparse 판정
    sheets = [{"columns": ["a", "b", "c", "", "", "", "", "", "", ""], "rows": []}]
    assert _is_sparse_sheets(sheets) is True


def _make_xlsx_with_empty_header(extra_cols: int) -> bytes:
    """헤더 첫 행이 거의 빈 채로 데이터 row 가 흩어진 PDF 인쇄형 .xlsx fixture."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("S")
    # 1행 = 헤더 (대부분 빈 문자열)
    ws.append([""] * extra_cols)
    # 2행 = 데이터지만 너무 흩어져서 표 의미 없음
    ws.append(["라벨", None, None, "값", None, None] + [None] * (extra_cols - 6))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_extract_uploaded_file_xlsx_falls_back_to_doc_text_when_sparse() -> None:
    """헤더가 모두 빈 문자열인 PDF 인쇄형 xlsx → tables=[] + doc_text 채워짐."""
    blob = _blob_with(_make_xlsx_with_empty_header(extra_cols=20))
    result = await extract_uploaded_file(blob, "p", "credit_report.xlsx")
    assert result is not None
    assert result["kind"] == "xlsx"
    assert result["content"] == []  # 표 추출 포기
    assert result["doc_text"] is not None
    assert "라벨" in result["doc_text"] and "값" in result["doc_text"]
    assert "=== 시트: S ===" in result["doc_text"]


async def test_extract_uploaded_file_xlsx_clean_sheet_keeps_tables_no_doc_text() -> None:
    """헤더가 정상인 깔끔한 xlsx → 기존대로 tables, doc_text=None."""
    blob = _blob_with(_make_xlsx({"재무": [["구분", "2023", "2024"], ["자산", 100, 200]]}))
    result = await extract_uploaded_file(blob, "p", "clean.xlsx")
    assert result is not None
    assert result["kind"] == "xlsx"
    assert len(result["content"]) == 1
    assert result["doc_text"] is None  # fallback 안 탐


async def test_extract_uploaded_file_xls_falls_back_to_doc_text_when_sparse() -> None:
    """레거시 .xls 도 동일 — 빈 헤더면 doc_text fallback."""
    # xlwt 로 셀 병합 흉내 — 빈 셀 가득한 헤더 + 흩어진 데이터
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Page 1")
    # 첫 행은 컬럼 31개 모두 빈 셀 (한국형 신용보고서 패턴)
    for c in range(31):
        ws.write(0, c, "")
    # 데이터는 라벨 + 값 한 쌍
    ws.write(1, 0, "상호")
    ws.write(1, 5, "테스트회사")
    buf = BytesIO()
    wb.save(buf)

    blob = _blob_with(buf.getvalue())
    result = await extract_uploaded_file(blob, "p", "report.xls")
    assert result is not None
    assert result["kind"] == "xlsx"
    assert result["content"] == []
    assert result["doc_text"] is not None
    assert "상호" in result["doc_text"] and "테스트회사" in result["doc_text"]


async def test_sheet_text_max_chars_truncates_per_sheet() -> None:
    """긴 시트는 SHEET_TEXT_MAX_CHARS 로 truncate (LLM 토큰 통제)."""
    long_text = "X" * (SHEET_TEXT_MAX_CHARS + 1000)
    # 빈 헤더 + 한 줄에 긴 텍스트 → sparse fallback 으로 들어감
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("S")
    ws.append([""] * 10)
    ws.append([long_text])
    buf = BytesIO()
    wb.save(buf)

    blob = _blob_with(buf.getvalue())
    result = await extract_uploaded_file(blob, "p", "huge.xlsx")
    assert result["doc_text"] is not None
    # 시트당 truncate marker 가 있어야 함
    assert "[...truncated]" in result["doc_text"]
    assert len(result["doc_text"]) <= SHEET_TEXT_MAX_CHARS + 100  # marker 약간 여유
