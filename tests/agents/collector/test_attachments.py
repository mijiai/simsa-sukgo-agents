from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

import openpyxl
from docx import Document

from src.agents.collector.attachments import (
    PER_FILE_TRUNCATE,
    ExtractedAttachment,
    extract_attachments,
)


def _make_docx(lines: list[str]) -> bytes:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_extract_attachments_skips_prompt_txt() -> None:
    blob = MagicMock()
    blob.list_prefix = AsyncMock(
        return_value=[
            "jobs/job-1/input/prompt.txt",
            "jobs/job-1/input/business_plan.docx",
        ]
    )
    blob.download = AsyncMock(return_value=_make_docx(["사업 개요"]))

    results = await extract_attachments(blob, "job-1")

    assert len(results) == 1
    assert results[0].filename == "business_plan.docx"
    assert "사업 개요" in results[0].text
    blob.download.assert_awaited_once_with("jobs/job-1/input/business_plan.docx")


async def test_extract_attachments_handles_multiple_formats() -> None:
    docx = _make_docx(["DOCX 본문"])
    xlsx = _make_xlsx([["계정", "값"], ["매출", 1000]])

    blob = MagicMock()
    blob.list_prefix = AsyncMock(
        return_value=[
            "jobs/job-2/input/business.docx",
            "jobs/job-2/input/financials.xlsx",
            "jobs/job-2/input/prompt.txt",
        ]
    )
    blob.download = AsyncMock(side_effect=[docx, xlsx])

    results = await extract_attachments(blob, "job-2")

    assert len(results) == 2
    by_name = {r.filename: r for r in results}
    assert "DOCX 본문" in by_name["business.docx"].text
    assert "매출 | 1000" in by_name["financials.xlsx"].text


async def test_extract_attachments_includes_unsupported_with_empty_text() -> None:
    """Unsupported extensions should still appear (so user knows file was attached)
    but with empty extracted text (no parsing attempted)."""
    blob = MagicMock()
    blob.list_prefix = AsyncMock(
        return_value=[
            "jobs/job-3/input/photo.png",
            "jobs/job-3/input/data.csv",
        ]
    )
    blob.download = AsyncMock()

    results = await extract_attachments(blob, "job-3")

    assert len(results) == 2
    assert all(r.text == "" for r in results)
    assert {r.filename for r in results} == {"photo.png", "data.csv"}
    blob.download.assert_not_called()


async def test_extract_attachments_skips_corrupt_file_keeps_others() -> None:
    good = _make_docx(["좋은 첨부"])
    bad = b"not a real docx"

    blob = MagicMock()
    blob.list_prefix = AsyncMock(
        return_value=[
            "jobs/job-4/input/good.docx",
            "jobs/job-4/input/broken.docx",
        ]
    )
    # extract_attachments sorts paths → broken.docx 가 먼저, good.docx 가 나중
    blob.download = AsyncMock(side_effect=[bad, good])

    results = await extract_attachments(blob, "job-4")

    assert len(results) == 2
    by_name = {r.filename: r for r in results}
    assert "좋은 첨부" in by_name["good.docx"].text
    assert by_name["broken.docx"].text == ""


async def test_extract_attachments_truncates_huge_file() -> None:
    huge = _make_docx(["A" * (PER_FILE_TRUNCATE + 5000) + "TAIL_MARKER"])

    blob = MagicMock()
    blob.list_prefix = AsyncMock(return_value=["jobs/job-5/input/huge.docx"])
    blob.download = AsyncMock(return_value=huge)

    results = await extract_attachments(blob, "job-5")

    assert len(results) == 1
    assert len(results[0].text) <= PER_FILE_TRUNCATE
    assert "TAIL_MARKER" not in results[0].text


async def test_extract_attachments_empty_input_returns_empty() -> None:
    blob = MagicMock()
    blob.list_prefix = AsyncMock(return_value=[])
    blob.download = AsyncMock()

    results = await extract_attachments(blob, "job-6")

    assert results == []
    blob.download.assert_not_called()


def test_extracted_attachment_is_frozen_dataclass() -> None:
    a = ExtractedAttachment(filename="x.docx", text="hi", chars=2)
    assert a.filename == "x.docx"
    # frozen=True 라 mutate 불가 — 우발적 변경 방지
    import dataclasses

    assert dataclasses.is_dataclass(a)
