from datetime import date, datetime
from io import BytesIO
from typing import Any

import openpyxl
import xlrd
from docx import Document
from pypdf import PdfReader
from xlrd import xldate_as_datetime

from src.config.logging import get_logger
from src.storage.blob_store import BlobStore

logger = get_logger(__name__)

_DOCX_SUFFIX = ".docx"
_PDF_SUFFIX = ".pdf"
_XLS_SUFFIX = ".xls"
_XLSX_SUFFIX = ".xlsx"
_SUPPORTED_SUFFIXES = (_DOCX_SUFFIX, _PDF_SUFFIX, _XLS_SUFFIX, _XLSX_SUFFIX)
_samples: list[str] = []


def extract_docx_text(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(text)
    return "\n".join(parts)


def _format_value(value: Any) -> str:
    """Stringify an Excel cell value, dropping float trailing .0 and ISO-formatting dates."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value).strip()


def extract_xlsx_text(data: bytes) -> str:
    """Read all sheets, serialize each as `=== 시트: <name> ===` + pipe-joined non-empty rows."""
    wb = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [_format_value(v) for v in row]
            cells = [c for c in cells if c]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"=== 시트: {sheet_name} ===\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)


def _xls_cell_value(cell: xlrd.sheet.Cell, datemode: int) -> Any:
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xldate_as_datetime(cell.value, datemode)
    return cell.value


def extract_xls_text(data: bytes) -> str:
    """Same shape as xlsx extractor but via xlrd (legacy .xls binary format)."""
    wb = xlrd.open_workbook(file_contents=data)
    parts: list[str] = []
    for sheet in wb.sheets():
        rows: list[str] = []
        for r in range(sheet.nrows):
            cells = [
                _format_value(_xls_cell_value(sheet.cell(r, c), wb.datemode))
                for c in range(sheet.ncols)
            ]
            cells = [c for c in cells if c]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"=== 시트: {sheet.name} ===\n" + "\n".join(rows))
    return "\n\n".join(parts)


def _extract(path: str, data: bytes) -> str:
    lower = path.lower()
    if lower.endswith(_DOCX_SUFFIX):
        return extract_docx_text(data)
    if lower.endswith(_PDF_SUFFIX):
        return extract_pdf_text(data)
    if lower.endswith(_XLSX_SUFFIX):
        return extract_xlsx_text(data)
    if lower.endswith(_XLS_SUFFIX):
        return extract_xls_text(data)
    return ""


async def load_financial_samples(blob: BlobStore, prefix: str) -> list[str]:
    """List + download + extract all .docx / .pdf / .xls / .xlsx samples under the prefix.

    Caches the result in module memory so the analyze service can read without I/O.
    Idempotent — calling again refreshes the cache.
    """
    global _samples
    paths = await blob.list_prefix(prefix)
    sample_paths = sorted(p for p in paths if p.lower().endswith(_SUPPORTED_SUFFIXES))

    samples: list[str] = []
    for path in sample_paths:
        try:
            data = await blob.download(path)
            text = _extract(path, data)
        except Exception as exc:
            logger.warning("financial.sample.load_failed", path=path, error=str(exc))
            continue
        if not text:
            logger.warning("financial.sample.empty", path=path)
            continue
        samples.append(text)
        logger.info("financial.sample.loaded", path=path, chars=len(text))

    _samples = samples
    logger.info("financial.samples.cached", count=len(samples), prefix=prefix)
    return samples


def get_cached_samples() -> list[str]:
    return list(_samples)


def reset_samples_for_tests(samples: list[str] | None = None) -> None:
    global _samples
    _samples = list(samples) if samples else []
