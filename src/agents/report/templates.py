from io import BytesIO

from docx import Document

from src.config.logging import get_logger
from src.storage.blob_store import BlobStore

logger = get_logger(__name__)

_DOCX_SUFFIX = ".docx"
_templates: list[str] = []


def extract_docx_text(data: bytes) -> str:
    """Extract paragraph + table text from a .docx byte stream."""
    doc = Document(BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


async def load_report_samples(blob: BlobStore, prefix: str) -> list[str]:
    """List + download + extract all .docx samples under the given Blob prefix.

    Caches the result in module memory so the report service can read without I/O.
    Idempotent — calling again refreshes the cache.
    """
    global _templates
    paths = await blob.list_prefix(prefix)
    docx_paths = sorted(p for p in paths if p.lower().endswith(_DOCX_SUFFIX))

    samples: list[str] = []
    for path in docx_paths:
        try:
            data = await blob.download(path)
            text = extract_docx_text(data)
        except Exception as exc:
            logger.warning("report.template.load_failed", path=path, error=str(exc))
            continue
        if not text:
            logger.warning("report.template.empty", path=path)
            continue
        samples.append(text)
        logger.info("report.template.loaded", path=path, chars=len(text))

    _templates = samples
    logger.info("report.templates.cached", count=len(samples), prefix=prefix)
    return samples


def get_cached_templates() -> list[str]:
    return list(_templates)


def reset_templates_for_tests(samples: list[str] | None = None) -> None:
    global _templates
    _templates = list(samples) if samples else []
