"""docx + markdown 렌더러 — SectionContent 를 두 가지 포맷으로 직접 빌드.

LLM 호출 없음. python-docx 로 .docx 바이트를, 별도 함수로 GFM markdown 을 생성.

docx 디자인:
- 섹션 헤딩: \"{number}. {title}\" Heading 1
- 표 제목: bold paragraph
- 표 본문: docx Table (헤더 1행 굵게, 본문 행)
- 출처 라벨: italic 8pt
- bullet: List Bullet style
- 이미지: BytesIO → add_picture(width=Cm(12)). 미제출 시 placeholder 텍스트

PR4 추가:
- base docx 스타일 상속 (Document(BytesIO(base)) + 본문 비우기)
- 필수 이미지 (optional=False) 누락 → 섹션 통째로 placeholder
- 별첨 섹션 자동 생성 (큰 표는 섹션 본문 대신 별첨에 게재)
"""

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt

from src.agents.report.sections import ImageContent, SectionContent, TableContent
from src.agents.report.template_spec import ReportTemplate
from src.common.constants import ReportSection
from src.config.logging import get_logger
from src.storage.blob_store import BlobStore

logger = get_logger(__name__)

_IMAGE_DEFAULT_WIDTH_CM = 12.0
_DATA_GAP_PLACEHOLDER = "[자료 미확보]"
_REQUIRED_IMAGE_MISSING_TEXT = "[자료 미확보 — 추가 제출 필요]"
_APPENDIX_TITLE = "별첨"
_BODY_REPLACED_NOTE = "(상세 표는 별첨 참조)"


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _open_base_document(base_bytes: bytes | None) -> Document:
    """base docx 가 주어지면 그 스타일/헤더/푸터를 상속하고 본문은 비운 새 Document 반환.

    base 가 None 이면 Document() 새로 생성.
    """
    if not base_bytes:
        return Document()
    try:
        doc = Document(BytesIO(base_bytes))
    except Exception as exc:
        logger.warning("report.render.base_open_failed", error=str(exc))
        return Document()
    # 본문 비우기 — body 의 paragraph/table 제거. sectPr (페이지 설정) 는 보존.
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)
    return doc


def _section_has_required_image_missing(section: SectionContent) -> bool:
    """필수 이미지(optional=False)가 data_gap 또는 blob_path 없음이면 True."""
    return any((not img.optional) and (img.data_gap or not img.blob_path) for img in section.images)


# ─────────────────────────── DOCX ───────────────────────────


def _add_table(doc: Document, table: TableContent) -> None:
    # 표 제목
    title_para = doc.add_paragraph()
    title_run = title_para.add_run(table.title)
    title_run.bold = True

    if table.data_gap or (not table.columns and not table.rows):
        gap_para = doc.add_paragraph()
        gap_run = gap_para.add_run(f"{_DATA_GAP_PLACEHOLDER} — 데이터 매핑 실패")
        gap_run.italic = True
    else:
        cols = table.columns or []
        rows = table.rows or []
        col_count = max(len(cols), max((len(r) for r in rows), default=0))
        if col_count == 0:
            col_count = 1
        docx_table = doc.add_table(rows=1 + len(rows), cols=col_count)
        docx_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        docx_table.style = "Table Grid"

        # 헤더 행
        header_cells = docx_table.rows[0].cells
        for i in range(col_count):
            cell_value = cols[i] if i < len(cols) else ""
            header_cells[i].text = _stringify(cell_value)
            for paragraph in header_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        # 본문 행
        for r_idx, row in enumerate(rows, start=1):
            row_cells = docx_table.rows[r_idx].cells
            for c_idx in range(col_count):
                cell_value = row[c_idx] if c_idx < len(row) else ""
                row_cells[c_idx].text = _stringify(cell_value)

    # 출처 라벨
    if table.source_label:
        src_para = doc.add_paragraph()
        src_run = src_para.add_run(table.source_label)
        src_run.italic = True
        src_run.font.size = Pt(8)

    # bullet 평가 (financial 이 작성한 것 그대로)
    for bullet in table.bullets:
        doc.add_paragraph(bullet, style="List Bullet")


async def _add_image(doc: Document, image: ImageContent, blob: BlobStore | None) -> None:
    title_para = doc.add_paragraph()
    title_para.add_run(image.title).bold = True

    if image.data_gap or not image.blob_path or blob is None:
        placeholder = doc.add_paragraph()
        placeholder_run = placeholder.add_run(
            f"{_DATA_GAP_PLACEHOLDER} — {image.suspected_role} 이미지 미제출"
        )
        placeholder_run.italic = True
        return

    try:
        data = await blob.download(image.blob_path)
        doc.add_picture(BytesIO(data), width=Cm(_IMAGE_DEFAULT_WIDTH_CM))
        if image.caption:
            cap_para = doc.add_paragraph()
            cap_run = cap_para.add_run(image.caption)
            cap_run.italic = True
            cap_run.font.size = Pt(9)
    except Exception as exc:
        logger.warning(
            "report.render.image_failed",
            blob_path=image.blob_path,
            error=str(exc),
        )
        err_para = doc.add_paragraph()
        err_run = err_para.add_run(f"{_DATA_GAP_PLACEHOLDER} — 이미지 로드 실패")
        err_run.italic = True


async def _render_section_docx(
    doc: Document,
    section: SectionContent,
    blob: BlobStore | None,
    appendix_table_ids: set[str] | None = None,
) -> None:
    """1개 섹션을 doc 에 append.

    필수 이미지가 빠져있으면 섹션 본문 대신 placeholder 만 출력.
    appendix_table_ids 에 들어있는 table_id 는 본문에서 생략하고 (별첨 참조) 안내만.
    """
    doc.add_heading(f"{section.number}. {section.title}", level=1)

    if _section_has_required_image_missing(section):
        # 필수 이미지 누락 → 섹션 통째로 placeholder
        msg = doc.add_paragraph()
        msg_run = msg.add_run(_REQUIRED_IMAGE_MISSING_TEXT)
        msg_run.italic = True
        return

    appendix_table_ids = appendix_table_ids or set()

    for table in section.tables:
        if table.table_id in appendix_table_ids:
            # 본문에는 표 제목 + "별첨 참조" 만 — 데이터는 별첨에서.
            title_para = doc.add_paragraph()
            title_para.add_run(table.title).bold = True
            ref_para = doc.add_paragraph()
            ref_run = ref_para.add_run(_BODY_REPLACED_NOTE)
            ref_run.italic = True
            # bullet 평가는 본문에 그대로
            for bullet in table.bullets:
                doc.add_paragraph(bullet, style="List Bullet")
        else:
            _add_table(doc, table)

    for image in section.images:
        await _add_image(doc, image, blob)

    # table_id 없는 일반 bullet 들 (예: 종합 의견 섹션)
    for bullet in section.section_bullets:
        doc.add_paragraph(bullet, style="List Bullet")


async def _render_appendix_docx(
    doc: Document, appendix_tables: list[TableContent], blob: BlobStore | None
) -> None:
    if not appendix_tables:
        return
    doc.add_heading(_APPENDIX_TITLE, level=1)
    for table in appendix_tables:
        _add_table(doc, table)


def collect_appendix_tables(
    sections: dict[ReportSection, SectionContent],
    *,
    row_threshold: int,
) -> tuple[list[TableContent], set[str]]:
    """row 가 threshold 초과인 표를 별첨 후보로 수집.

    Returns:
        (appendix_tables, appendix_table_ids)
        - appendix_tables: 별첨 섹션에 게재할 TableContent 목록 (전체 데이터 보존)
        - appendix_table_ids: 본문에서 생략 처리할 table_id set
    """
    appendix: list[TableContent] = []
    ids: set[str] = set()
    for section in sections.values():
        for table in section.tables:
            if table.data_gap:
                continue
            if len(table.rows) > row_threshold:
                appendix.append(table)
                ids.add(table.table_id)
    return appendix, ids


async def render_report_docx(
    template: ReportTemplate,
    sections: dict[ReportSection, SectionContent],
    blob: BlobStore | None = None,
    *,
    base_docx_bytes: bytes | None = None,
    appendix_row_threshold: int | None = None,
) -> bytes:
    """전체 보고서를 docx 바이트로 렌더.

    blob 가 None 이면 이미지는 모두 placeholder 처리 (테스트/오프라인 모드).
    base_docx_bytes 가 주어지면 그 docx 의 스타일/헤더/푸터를 상속.
    appendix_row_threshold 초과 행 수의 표는 본문에서 빼고 별첨 섹션에 게재.
    """
    doc = _open_base_document(base_docx_bytes)

    # 표지 — 단순한 제목만 (base 에 표지가 있어도 덮어쓰지 않고 추가 — base body 는 비워진 상태)
    title = doc.add_paragraph()
    title_run = title.add_run("기업 여신 심사 보고서")
    title_run.bold = True
    title_run.font.size = Pt(20)

    appendix_tables: list[TableContent] = []
    appendix_ids: set[str] = set()
    if appendix_row_threshold and appendix_row_threshold > 0:
        appendix_tables, appendix_ids = collect_appendix_tables(
            sections, row_threshold=appendix_row_threshold
        )

    for spec in template.sections:
        section = sections.get(spec.section_id)
        if section is None:
            # spec 에는 있지만 sections 에 없는 경우 — placeholder 만 출력
            doc.add_heading(f"{spec.number}. {spec.title}", level=1)
            err = doc.add_paragraph()
            err.add_run(f"{_DATA_GAP_PLACEHOLDER} — 섹션 데이터 없음").italic = True
            continue
        await _render_section_docx(doc, section, blob, appendix_table_ids=appendix_ids)

    await _render_appendix_docx(doc, appendix_tables, blob)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─────────────────────────── Markdown ───────────────────────────


def _md_table(table: TableContent) -> list[str]:
    lines: list[str] = [f"**{table.title}**", ""]
    if table.data_gap or (not table.columns and not table.rows):
        lines.append(f"_{_DATA_GAP_PLACEHOLDER} — 데이터 매핑 실패_")
        lines.append("")
        return lines
    cols = table.columns or []
    rows = table.rows or []
    col_count = max(len(cols), max((len(r) for r in rows), default=0))
    if col_count == 0:
        col_count = 1
    header = (
        "| "
        + " | ".join(_stringify(cols[i] if i < len(cols) else "") for i in range(col_count))
        + " |"
    )
    sep = "| " + " | ".join("---" for _ in range(col_count)) + " |"
    lines.append(header)
    lines.append(sep)
    for row in rows:
        lines.append(
            "| "
            + " | ".join(_stringify(row[i] if i < len(row) else "") for i in range(col_count))
            + " |"
        )
    lines.append("")
    if table.source_label:
        lines.append(f"_{table.source_label}_")
        lines.append("")
    for bullet in table.bullets:
        lines.append(f"- {bullet}")
    if table.bullets:
        lines.append("")
    return lines


def _md_image(image: ImageContent) -> list[str]:
    lines = [f"**{image.title}**", ""]
    if image.data_gap or not image.blob_path:
        lines.append(f"_{_DATA_GAP_PLACEHOLDER} — {image.suspected_role} 이미지 미제출_")
        lines.append("")
        return lines
    # markdown 은 SAS URL 이 없으니 blob_path 만 placeholder 로 표기
    lines.append(f"![{image.title}]({image.blob_path})")
    if image.caption:
        lines.append(f"_{image.caption}_")
    lines.append("")
    return lines


def _md_section(
    section: SectionContent,
    appendix_ids: set[str] | None = None,
) -> list[str]:
    lines = [f"# {section.number}. {section.title}", ""]
    if _section_has_required_image_missing(section):
        lines.append(f"_{_REQUIRED_IMAGE_MISSING_TEXT}_")
        lines.append("")
        return lines
    appendix_ids = appendix_ids or set()
    for table in section.tables:
        if table.table_id in appendix_ids:
            lines.append(f"**{table.title}**")
            lines.append("")
            lines.append(f"_{_BODY_REPLACED_NOTE}_")
            lines.append("")
            for bullet in table.bullets:
                lines.append(f"- {bullet}")
            if table.bullets:
                lines.append("")
        else:
            lines.extend(_md_table(table))
    for image in section.images:
        lines.extend(_md_image(image))
    for bullet in section.section_bullets:
        lines.append(f"- {bullet}")
    if section.section_bullets:
        lines.append("")
    return lines


def render_report_markdown(
    template: ReportTemplate,
    sections: dict[ReportSection, SectionContent],
    *,
    appendix_row_threshold: int | None = None,
) -> str:
    """전체 보고서를 GFM markdown 으로 렌더 (호환용)."""
    out: list[str] = ["# 기업 여신 심사 보고서", ""]
    appendix_tables: list[TableContent] = []
    appendix_ids: set[str] = set()
    if appendix_row_threshold and appendix_row_threshold > 0:
        appendix_tables, appendix_ids = collect_appendix_tables(
            sections, row_threshold=appendix_row_threshold
        )
    for spec in template.sections:
        section = sections.get(spec.section_id)
        if section is None:
            out.append(f"# {spec.number}. {spec.title}")
            out.append("")
            out.append(f"_{_DATA_GAP_PLACEHOLDER} — 섹션 데이터 없음_")
            out.append("")
            continue
        out.extend(_md_section(section, appendix_ids=appendix_ids))
    if appendix_tables:
        out.append(f"# {_APPENDIX_TITLE}")
        out.append("")
        for table in appendix_tables:
            out.extend(_md_table(table))
    return "\n".join(out)
